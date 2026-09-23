#!/usr/bin/env python3
"""
Arma el registro de cotizaciones (Excel/Google Sheet) desde cero.

Existe porque el registro ya se perdió una vez: se generó, se entregó, y el archivo no quedó
en ninguna parte versionada — y todo el sistema depende de él (ahí vive el correlativo, los
costos reales con los que se va a deducir la fórmula de precio, y el conteo de subtipos del
que depende la promoción a tipo propio). El archivo en sí NO se versiona (es dato operativo
que cambia todo el tiempo, ver criterios-actuales.md punto 6); lo que se versiona es este
script que lo reconstruye.

Los totales, la ganancia y el margen son FÓRMULAS, nunca valores escritos: si se escriben a
mano dejan de cuadrar en cuanto alguien corrige un costo.

Uso:
    python3 herramientas/armar_registro.py [salida.xlsx]
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

VERDE = "1E7B4F"
AMARILLO = "FFF59D"
GRIS = "EFEFEF"
NARANJA = "E8833A"

# (encabezado, ancho, etapa). Etapa 1 = al cotizar, 2 = cuando el cliente responde,
# 3 = al terminar el trabajo, F = fórmula (no se escribe a mano).
COLUMNAS = [
    ("N° Cotización",          15, 1),
    ("Fecha cotización",       16, 1),
    ("Cliente",                26, 1),
    ("Local",                  16, 1),
    ("Tipo de servicio",       18, 1),
    ("Subtipo",                18, 1),
    ("Descripción breve",      40, 1),
    ("Unidades (equipos)",     13, 1),
    ("Subtotal",               12, 1),
    ("IGV",                    11, 1),
    ("Total",                  12, 1),
    ("Estado",                 14, 2),
    ("Motivo de pérdida",      24, 2),
    ("Vueltas de corrección",  18, 2),
    ("Fecha fin trabajo",      16, 3),
    ("Personas",               10, 3),
    ("Amanecidas",             11, 3),
    ("Costo repuestos",        15, 3),
    ("Costo mano de obra",     17, 3),
    ("Costo logística",        15, 3),
    ("Costo total",            13, "F"),
    ("Ganancia",               12, "F"),
    ("Markup % (s/costo)",     13, "F"),
    ("Unid. por amanecida",    17, "F"),
    ("Precio por unidad",      15, "F"),
    # Agregadas el 23/09/2026 AL FINAL, a propósito: el registro vivo es un Google Sheet, y
    # agregar columnas a la derecha no mueve ninguna fórmula de las que ya existen.
    # Sin costo estimado ni precio calculado, el ajuste final por experiencia era invisible y
    # las cotizaciones perdidas nunca tenían costo (hallazgos-abiertos.md, N05 y N06).
    ("Costo estimado",         15, 1),
    ("Subtotal calculado",     17, 1),
    ("Nota interna",           40, 1),
    ("Ajuste %",               11, "F"),
    ("Markup cotizado %",      16, "F"),
    ("Margen % (s/venta)",     16, "F"),
]


def letra(nombre):
    """Letra de la columna por su encabezado: las fórmulas nuevas no dependen de contar a mano."""
    return get_column_letter([c[0] for c in COLUMNAS].index(nombre) + 1)

COLOR_ETAPA = {1: VERDE, 2: NARANJA, 3: "5B6770", "F": "8E44AD"}

ESTADOS = "Cotizado,Enviado,Ganado,Perdido,Sin respuesta"
MOTIVOS = "Precio alto,Plazo,Otro proveedor,Sin presupuesto,No respondió,Otro"
# Tipos y subtipos salen de servicios.json, la fuente única. Antes estaban copiados a mano acá,
# y agregar un subtipo allá no lo agregaba al desplegable (hallazgos-abiertos.md, N32).
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SERVICIOS = json.load(open(os.path.join(RAIZ, "servicios.json"), encoding="utf-8"))
TIPOS = ",".join(dict.fromkeys(t["categoria_registro"] for t in SERVICIOS["tipos"].values()))
SUBTIPOS = ",".join(SERVICIOS["subtipos_sugeridos"])

FILAS = 400          # espacio para fórmulas por adelantado
PRIMERA = 3          # fila 1 = leyenda de etapas, fila 2 = encabezados


def hoja_cotizaciones(wb, semilla):
    ws = wb.active
    ws.title = "Cotizaciones"

    # Fila 1: de qué etapa es cada bloque de columnas. Sin esto nadie sabe qué se llena cuándo.
    for i, (_, ancho, etapa) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=1, column=i)
        c.value = {1: "1 · al cotizar", 2: "2 · cuando responde",
                   3: "3 · al terminar", "F": "se calcula solo"}[etapa]
        c.font = Font(size=8, italic=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COLOR_ETAPA[etapa])
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho

    borde = Border(bottom=Side(style="thin", color="CCCCCC"))
    for i, (nombre, _, etapa) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=2, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=COLOR_ETAPA[etapa])
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borde
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "B3"

    c_est, s_calc = letra("Costo estimado"), letra("Subtotal calculado")
    aj, mk, mg = letra("Ajuste %"), letra("Markup cotizado %"), letra("Margen % (s/venta)")
    for f in range(PRIMERA, PRIMERA + FILAS):
        # Ajuste final por experiencia: cuánto se movió el precio respecto del calculado.
        ws[f"{aj}{f}"] = f'=IF(OR({s_calc}{f}="",{s_calc}{f}=0,I{f}=""),"",I{f}/{s_calc}{f}-1)'
        # Markup al cotizar: existe para TODAS las cotizaciones, se ganen o se pierdan. Es lo que
        # permite ver el techo de precio comparando ganadas contra perdidas.
        ws[f"{mk}{f}"] = f'=IF(OR({c_est}{f}="",{c_est}{f}=0,I{f}=""),"",I{f}/{c_est}{f}-1)'
        # Margen sobre venta, el que usa la contabilidad: un markup de 73% es un margen de 42%.
        ws[f"{mg}{f}"] = f'=IF(OR(V{f}="",I{f}="",I{f}=0),"",V{f}/I{f})'
        for c in (c_est, s_calc):
            ws[f"{c}{f}"].number_format = '#,##0.00'
        for c in (aj, mk, mg):
            ws[f"{c}{f}"].number_format = '0.0%'
            ws[f"{c}{f}"].fill = PatternFill("solid", fgColor=GRIS)

        # Costo total = repuestos + mano de obra + logística
        ws[f"U{f}"] = f"=IF(COUNT(R{f}:T{f})=0,\"\",SUM(R{f}:T{f}))"
        # Ganancia sobre el SUBTOTAL: el IGV no es ingreso, se le entrega a SUNAT.
        # Contarlo como ganancia la infla ~36% (criterios-actuales.md punto 6).
        ws[f"V{f}"] = f"=IF(OR(U{f}=\"\",I{f}=\"\"),\"\",I{f}-U{f})"
        ws[f"W{f}"] = f"=IF(OR(U{f}=\"\",U{f}=0),\"\",V{f}/U{f})"
        ws[f"X{f}"] = f"=IF(OR(Q{f}=\"\",Q{f}=0,H{f}=\"\"),\"\",H{f}/Q{f})"
        ws[f"Y{f}"] = f"=IF(OR(H{f}=\"\",H{f}=0,I{f}=\"\"),\"\",I{f}/H{f})"
        for col in "IJKRSTUVY":
            ws[f"{col}{f}"].number_format = '#,##0.00'
        ws[f"W{f}"].number_format = '0.0%'
        ws[f"X{f}"].number_format = '0.0'
        for col in "UVWXY":
            ws[f"{col}{f}"].fill = PatternFill("solid", fgColor=GRIS)

    rango = f"{PRIMERA}:{PRIMERA + FILAS - 1}"
    for col, lista in (("L", ESTADOS), ("M", MOTIVOS), ("E", TIPOS), ("F", SUBTIPOS)):
        dv = DataValidation(type="list", formula1=f'"{lista}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}{PRIMERA}:{col}{PRIMERA + FILAS - 1}")

    for j, valor in enumerate(semilla, start=1):
        ws.cell(row=PRIMERA, column=j, value=valor)
    return ws


def hoja_resumen(wb):
    ws = wb.create_sheet("Resumen")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    d = f"Cotizaciones!"
    r = f"{PRIMERA}:{PRIMERA + FILAS - 1}"

    def rng(columna):
        return f"{d}{columna}{PRIMERA}:{columna}{PRIMERA + FILAS - 1}"

    bloques = [
        ("COMERCIAL", [
            ("Cotizaciones emitidas", f'=COUNTA({d}A{PRIMERA}:A{PRIMERA+FILAS-1})'),
            ("Total cotizado (subtotal)", f'=SUM({d}I{PRIMERA}:I{PRIMERA+FILAS-1})'),
            ("Ganadas", f'=COUNTIF({d}L{PRIMERA}:L{PRIMERA+FILAS-1},"Ganado")'),
            ("Perdidas", f'=COUNTIF({d}L{PRIMERA}:L{PRIMERA+FILAS-1},"Perdido")'),
            # Autocontenida a propósito: referenciar B3/B5 la rompía en silencio al mover una
            # fila del resumen (pasó: dividía perdidas entre el monto cotizado).
            ("% ganadas",
             f'=IF(COUNTA({d}A{PRIMERA}:A{PRIMERA+FILAS-1})=0,"",'
             f'COUNTIF({d}L{PRIMERA}:L{PRIMERA+FILAS-1},"Ganado")/'
             f'COUNTA({d}A{PRIMERA}:A{PRIMERA+FILAS-1}))'),
            ("Facturado (ganadas, subtotal)",
             f'=SUMIF({d}L{PRIMERA}:L{PRIMERA+FILAS-1},"Ganado",{d}I{PRIMERA}:I{PRIMERA+FILAS-1})'),
        ]),
        # "Markup" y no "margen": es ganancia ÷ costo. El margen contable es ganancia ÷ venta, y
        # confundirlos hace leer 73% donde el margen real es 42% (hallazgos-abiertos.md, N11).
        ("RENTABILIDAD REAL POR TIPO", [
            ("Markup medio (s/costo) — Carretillas",
             f'=IFERROR(AVERAGEIF({d}E{PRIMERA}:E{PRIMERA+FILAS-1},"Carretillas",{d}W{PRIMERA}:W{PRIMERA+FILAS-1}),"sin datos")'),
            ("Markup medio (s/costo) — Balanzas",
             f'=IFERROR(AVERAGEIF({d}E{PRIMERA}:E{PRIMERA+FILAS-1},"Balanzas",{d}W{PRIMERA}:W{PRIMERA+FILAS-1}),"sin datos")'),
            ("Markup medio (s/costo) — Servicios Generales",
             f'=IFERROR(AVERAGEIF({d}E{PRIMERA}:E{PRIMERA+FILAS-1},"Servicios Generales",{d}W{PRIMERA}:W{PRIMERA+FILAS-1}),"sin datos")'),
            ("Margen medio (s/venta) — todos los tipos",
             f'=IFERROR(AVERAGE({rng(letra("Margen % (s/venta)"))}),"sin datos")'),
        ]),
        # Lo que el precio calculado no explica. Se lee junto con el resultado: si las ganadas
        # y las perdidas se cotizaron con el mismo markup, el precio no es lo que decide.
        ("PRECIO: CALCULADO vs FINAL", [
            ("Ajuste medio sobre el precio calculado",
             f'=IFERROR(AVERAGE({rng(letra("Ajuste %"))}),"sin datos")'),
            ("Markup cotizado medio — ganadas",
             f'=IFERROR(AVERAGEIF({rng("L")},"Ganado",{rng(letra("Markup cotizado %"))}),"sin datos")'),
            ("Markup cotizado medio — perdidas",
             f'=IFERROR(AVERAGEIF({rng("L")},"Perdido",{rng(letra("Markup cotizado %"))}),"sin datos")'),
            ("Markup cotizado medio — perdidas por precio alto",
             f'=IFERROR(AVERAGEIFS({rng(letra("Markup cotizado %"))},{rng("L")},"Perdido",'
             f'{rng("M")},"Precio alto"),"sin datos")'),
            ("⚠ Cotizaciones sin costo estimado",
             f'=COUNTIFS({rng("A")},"<>",{rng(letra("Costo estimado"))},"")'),
        ]),
        ("CALIBRACIÓN DE LA FÓRMULA DE PRECIO — se lee al cierre de mes", [
            ("Unidades por amanecida (promedio real)",
             f'=IFERROR(AVERAGE({d}X{PRIMERA}:X{PRIMERA+FILAS-1}),"sin datos")'),
            ("Costo real de mano de obra por amanecida",
             f'=IFERROR(SUM({d}S{PRIMERA}:S{PRIMERA+FILAS-1})/SUM({d}Q{PRIMERA}:Q{PRIMERA+FILAS-1}),"sin datos")'),
            ("Precio promedio por unidad",
             f'=IFERROR(AVERAGE({d}Y{PRIMERA}:Y{PRIMERA+FILAS-1}),"sin datos")'),
            ("⚠ Trabajos ganados sin personas/amanecidas",
             f'=SUMPRODUCT(({d}L{PRIMERA}:L{PRIMERA+FILAS-1}="Ganado")*'
             f'(({d}P{PRIMERA}:P{PRIMERA+FILAS-1}="")+({d}Q{PRIMERA}:Q{PRIMERA+FILAS-1}="")>0))'),
        ]),
        ("CALIDAD", [
            ("Vueltas de corrección (promedio)",
             f'=IFERROR(AVERAGE({d}N{PRIMERA}:N{PRIMERA+FILAS-1}),"sin datos")'),
            ("Cotizaciones que salieron sin corrección",
             f'=COUNTIF({d}N{PRIMERA}:N{PRIMERA+FILAS-1},0)'),
        ]),
        ("PROMOCIÓN DE SUBTIPO — umbral: 3 del mismo subtipo", []),
    ]

    fila = 1
    for titulo, filas in bloques:
        c = ws.cell(row=fila, column=1, value=titulo)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=VERDE)
        ws.cell(row=fila, column=2).fill = PatternFill("solid", fgColor=VERDE)
        fila += 1
        for etq, formula in filas:
            ws.cell(row=fila, column=1, value=etq)
            cel = ws.cell(row=fila, column=2, value=formula)
            if "%" in etq or any(p in etq for p in ("Margen", "Markup", "Ajuste")):
                cel.number_format = '0.0%'
            elif any(p in etq for p in ("cotizado", "Facturado", "Costo", "Precio")):
                cel.number_format = '#,##0.00'
            cel.fill = PatternFill("solid", fgColor=AMARILLO)
            fila += 1
        fila += 1

    for sub in SUBTIPOS.split(","):
        ws.cell(row=fila, column=1, value=f"  {sub}")
        cel = ws.cell(row=fila, column=2,
                      value=f'=COUNTIF({d}F{PRIMERA}:F{PRIMERA+FILAS-1},"{sub}")')
        cel.fill = PatternFill("solid", fgColor=AMARILLO)
        ws.cell(row=fila, column=3,
                value=f'=IF(B{fila}>=3,"⬅ evaluar tipo propio al cierre de mes","")')
        fila += 1
    return ws


def hoja_instrucciones(wb):
    ws = wb.create_sheet("Cómo se llena")
    ws.column_dimensions["A"].width = 100
    texto = [
        ("REGISTRO DE COTIZACIONES — cómo se llena", True),
        ("", False),
        ("Se llena en 3 momentos, no de una vez. Los colores de la fila 1 dicen cuál es cuál.", False),
        ("", False),
        ("1) AL COTIZAR (verde) — la fila la entrega el generador ya armada, se pega y listo:", True),
        ("   N°, fecha, cliente, local, tipo, subtipo, descripción, unidades, subtotal, IGV, total.", False),
        ("   El generador imprime esa línea al final; no hay que transcribir nada a mano.", False),
        ("   UNIDADES = equipos atendidos (10 carretillas), NO la suma de repuestos. Sumar ruedas y", False),
        ("   botellas daba 33 en vez de 10 y descuadraba el precio por unidad.", False),
        ("   Al final de la hoja van tres más de esta etapa: COSTO ESTIMADO, SUBTOTAL CALCULADO y", False),
        ("   NOTA INTERNA. Se llenan en TODAS las cotizaciones, se ganen o se pierdan: sin eso no se", False),
        ("   puede medir cuánto se ajusta el precio por experiencia, ni ver a qué markup se pierde.", False),
        ("", False),
        ("2) CUANDO EL CLIENTE RESPONDE (naranja):", True),
        ("   Estado, y si se perdió, EL MOTIVO. Sin el motivo, el % de perdidas no dice qué cambiar.", False),
        ("   Vueltas de corrección: cuántas veces se pidió corregir antes de aprobarla. 0 si salió", False),
        ("   derecho. Es el único dato que mide calidad — no se puede sacar de ningún log.", False),
        ("", False),
        ("3) AL TERMINAR EL TRABAJO (gris), solo si se ganó:", True),
        ("   Fecha fin, PERSONAS, AMANECIDAS y los tres costos reales.", False),
        ("   Personas y amanecidas son las dos casillas más importantes del archivo: son", False),
        ("   exactamente lo que falta para calcular el precio en vez de estimarlo. Si quedan", False),
        ("   vacías, al cierre del mes habrá facturación pero no se podrá deducir la fórmula.", False),
        ("", False),
        ("LO MORADO NO SE TOCA: costo total, ganancia, markup, unidades por amanecida, precio por", True),
        ("unidad, ajuste %, markup cotizado y margen son fórmulas. Escribirlas a mano las descuadra", False),
        ("en cuanto se corrija un costo.", False),
        ("", False),
        ("MARKUP no es MARGEN. Markup = ganancia ÷ costo. Margen = ganancia ÷ venta. Un markup de", True),
        ("73% es un margen de 42%. La contabilidad y cualquier referencia del rubro usan margen.", False),
        ("", False),
        ("La ganancia se calcula sobre el SUBTOTAL, no sobre el total: el IGV no es ingreso de la", True),
        ("empresa, se cobra al cliente y se entrega a SUNAT. Contarlo infla la ganancia ~36%.", False),
        ("", False),
        ("Para subirlo a Google Sheets: Archivo > Importar > este .xlsx. Las listas desplegables y", False),
        ("las fórmulas se mantienen.", False),
    ]
    for i, (linea, negrita) in enumerate(texto, start=1):
        c = ws.cell(row=i, column=1, value=linea)
        if negrita:
            c.font = Font(bold=True)
    return ws


if __name__ == "__main__":
    salida = sys.argv[1] if len(sys.argv) > 1 else "Registro_Cotizaciones_MunayPeru.xlsx"
    # Primera fila ya cargada: la única cotización emitida a la fecha.
    semilla = ["COT-2026-001", "17/09/2026", "Hipermercados Tottus S.A", "Atocongo",
               "Servicios Generales", "laminado-seguridad",
               "Cambio de laminado de seguridad certificado en espejos y vidrio templado",
               63, 4980.00, 896.40, 5876.40, "Enviado"]
    wb = Workbook()
    hoja_cotizaciones(wb, semilla)
    hoja_resumen(wb)
    hoja_instrucciones(wb)
    wb.save(salida)
    print(f"Registro armado: {salida}")
    print(f"  Hoja 'Cotizaciones': {len(COLUMNAS)} columnas, {FILAS} filas con fórmulas listas")
    print("  Hoja 'Resumen': KPIs comerciales, rentabilidad, calibración, calidad y conteo de subtipos")
    print("  Hoja 'Cómo se llena': las 3 etapas explicadas")
    print(f"  Fila 1 cargada: {semilla[0]} — {semilla[2]}")
    nuevas = COLUMNAS[[c[0] for c in COLUMNAS].index("Costo estimado"):]
    print("\nSi el registro ya existe en Google Sheets, no hace falta rehacerlo: agregar al final")
    print("estas columnas, en este orden (ninguna fórmula existente se mueve):")
    for nombre, _, etapa in nuevas:
        print(f"  {letra(nombre)}  {nombre}{'  (fórmula: copiarla de este archivo)' if etapa == 'F' else ''}")
    print("Y renombrar H → 'Unidades (equipos)' y W → 'Markup % (s/costo)' (mismas fórmulas).")
