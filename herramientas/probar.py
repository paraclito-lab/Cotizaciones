#!/usr/bin/env python3
"""
Pruebas de regresión del sistema de cotizaciones. Correrlas antes de subir cualquier cambio:

    python3 herramientas/probar.py            rápidas, sin navegador (segundos)
    python3 herramientas/probar.py --render   además genera PNG y PDF reales con Chromium

La referencia son cotizaciones REALES ya emitidas: si un cambio mueve un céntimo de lo que ya
se le entregó a un cliente, falla. Agregar un trabajo nuevo a la biblioteca no debería romper
ninguna — para eso existen.
"""
import datetime, json, os, shutil, subprocess, sys, tempfile, traceback
from decimal import Decimal

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
sys.path.insert(0, os.path.join(RAIZ, "plantilla"))
import generar  # noqa: E402
import cotizar_trabajo as motor  # noqa: E402

PRUEBAS = []
DIA = datetime.date(2026, 9, 22)


class Omitida(Exception):
    pass


def prueba(lenta=False):
    def registrar(f):
        f.lenta = lenta
        PRUEBAS.append(f)
        return f
    return registrar


def cargar(ruta):
    return json.load(open(os.path.join(RAIZ, ruta), encoding="utf-8"))


def validar(cot, **modos):
    servicio = generar.tipo_de(cot)
    sub, _, _ = generar.calcular(cot)
    return generar.validar(cot, servicio, sub, **modos)


def igual(obtenido, esperado, que=""):
    assert obtenido == esperado, f"{que}: se esperaba {esperado!r}, salió {obtenido!r}"


def contiene(textos, fragmento, que=""):
    assert any(fragmento in t for t in textos), f"{que}: ningún mensaje contiene {fragmento!r}: {textos}"


EMITIDAS = {
    "cotizaciones-emitidas/9196-2026.json": ("6400", "1152", "7552"),
    "cotizaciones-emitidas/9197-2026.json": ("2318", "417.24", "2735.24"),
    "herramientas/ejemplo-trabajo-cotizacion.json": ("4980", "896.40", "5876.40"),
}


# ---------------------------------------------------------------- generador

@prueba()
def totales_de_cotizaciones_emitidas_no_cambian():
    for ruta, esperado in EMITIDAS.items():
        igual(generar.calcular(cargar(ruta)), tuple(Decimal(x) for x in esperado), ruta)


@prueba()
def cotizaciones_emitidas_pasan_la_validacion():
    for ruta in EMITIDAS:
        errores, alertas, _ = validar(cargar(ruta))
        igual(errores, [], ruta)
        assert not any("correlativo" in a for a in alertas), f"{ruta}: alerta de N° falsa: {alertas}"


@prueba()
def piso_de_carretillas_avisa_sin_bloquear():
    _, alertas, _ = validar(cargar("cotizaciones-emitidas/9197-2026.json"))
    contiene(alertas, "bajo el piso", "9197 a S/230")
    _, alertas, _ = validar(cargar("cotizaciones-emitidas/9196-2026.json"))
    assert not any("piso" in a for a in alertas), "9196 a S/500 no debería avisar piso"


@prueba()
def igv_con_redondeo_contable():
    for subtotal, igv in (("30.25", "5.45"), ("2.25", "0.41"), ("102.25", "18.41"), ("12.25", "2.21")):
        cot = {"numero": "1-2026", "tipo": "carretillas", "cliente": "X", "fecha": "hoy",
               "items": [{"cant": 1, "descripcion": "x", "precio_unit": float(subtotal)}]}
        igual(generar.calcular(cot)[1], Decimal(igv), f"IGV de {subtotal}")


@prueba()
def unidades_cuenta_equipos_no_repuestos():
    cot = cargar("cotizaciones-emitidas/9196-2026.json")
    igual(generar.unidades_de(cot), 10, "explícitas")
    del cot["unidades"]
    igual(generar.unidades_de(cot), 10, "sin 'unidades' toma el ítem 1, no la suma (33)")
    _, _, avisos = validar(cot)
    contiene(avisos, "sin 'unidades'", "aviso de unidades inferidas")


@prueba()
def texto_interno_no_llega_al_cliente():
    for campo, texto in (("nota", "Precio según catalogo.md"),
                         ("nota", "Bajo el piso mínimo por excepción"),
                         ("condiciones", ["Markup reducido por volumen"])):
        cot = cargar("cotizaciones-emitidas/9196-2026.json")
        cot[campo] = texto
        errores, _, _ = validar(cot)
        contiene(errores, "texto interno", f"{campo}={texto!r}")
    cot = cargar("cotizaciones-emitidas/9196-2026.json")
    cot["items"][0]["descripcion"] += " Por validar."
    contiene(validar(cot)[0], "texto interno", "descripción de ítem")
    cot = cargar("cotizaciones-emitidas/9196-2026.json")
    igual(validar(cot)[0], [], "nota_interna sí puede hablar de catálogo y piso")


@prueba()
def final_exige_numero_real_y_precios_completos():
    cot = cargar("cotizaciones-emitidas/9196-2026.json")
    igual(validar(cot, final=True, archivar_pedido=True)[0], [], "9196 lista para --final")
    cot["numero"] = "PENDIENTE-2026"
    contiene(validar(cot, final=True, archivar_pedido=True)[0], "--final exige", "N° pendiente")
    cot = cargar("cotizaciones-emitidas/9196-2026.json")
    del cot["items"][1]["precio_unit"]
    cot["items"][1]["pendiente"] = True
    contiene(validar(cot, final=True, archivar_pedido=True)[0], "precio pendiente", "ítem pendiente")


@prueba()
def sin_precios_no_muestra_ceros_ni_alertas_repetidas():
    cot = {"numero": "PENDIENTE-2026", "tipo": "generales", "subtipo": "otros", "cliente": "X",
           "fecha": "hoy", "items": [{"cant": 3, "descripcion": "a", "pendiente": True},
                                     {"cant": 1, "descripcion": "b", "pendiente": True}]}
    errores, alertas, _ = validar(cot, sin_precios=True)
    igual(errores, [], "propuesta técnica válida")
    contiene(alertas, "SIN PRECIOS", "alerta general")
    assert not any("ítem" in a and "PENDIENTE" in a for a in alertas), f"alertas repetidas: {alertas}"
    contiene(validar(cot, sin_precios=True, archivar_pedido=True)[0], "no se archiva", "no archivable")
    cot["items"][0] = {"cant": 1, "descripcion": "a", "precio_unit": 10}
    contiene(validar(cot, sin_precios=True)[0], "--sin-precios pero hay", "con precios")
    html = generar.construir_html(dict(cot, items=[{"cant": 1, "descripcion": "a", "pendiente": True}]),
                                  generar.tipo_de(cot), Decimal(0), Decimal(0), Decimal(0), True)
    celdas = html[html.index("<tfoot"):html.index("</tfoot>")]   # el CSS también dice "S/ 0.00"
    assert "Por definir" in celdas and "S/ 0.00" not in celdas, "los totales deben decir Por definir"


@prueba()
def nombre_del_pdf_de_entrega():
    for ruta, esperado in (
            ("cotizaciones-emitidas/9196-2026.json", "N° 9196-2026 CARRETILLAS TT PUCALLPA MUNAY.pdf"),
            ("cotizaciones-emitidas/9197-2026.json",
             "N° 9197-2026 CARRETILLAS HB PRECIO UNO PANAMERICANA URBANA MUNAY.pdf")):
        cot = cargar(ruta)
        igual(generar.nombre_entrega(cot, generar.tipo_de(cot)), esperado, ruta)


@prueba()
def recorte_en_palabra_completa():
    largo = "Mantenimiento preventivo de 7 carretillas hidráulicas (stockas N° 1 al 7): cambio"
    igual(generar.recortar(largo), "Mantenimiento preventivo de 7 carretillas hidráulicas…")
    igual(generar.recortar("corto"), "corto")


@prueba()
def plantilla_sin_marcadores_sueltos_y_atencion_opcional():
    for ruta in EMITIDAS:
        cot = cargar(ruta)
        s, i, t = generar.calcular(cot)
        html = generar.construir_html(cot, generar.tipo_de(cot), s, i, t)
        assert "{{" not in html, f"{ruta}: quedó un marcador sin llenar"
        tiene_atencion = "Atenci&oacute;n:" in html
        igual(tiene_atencion, bool(cot.get("atencion", "").strip()), f"{ruta}: línea de Atención")
    cot = dict(cargar("cotizaciones-emitidas/9197-2026.json"), condiciones=["Una condición."])
    s, i, t = generar.calcular(cot)
    assert "Condiciones:" in generar.construir_html(cot, generar.tipo_de(cot), s, i, t)


# ---------------------------------------------------------------- catálogo y fichas

@prueba()
def catalogo_se_lee_y_no_adivina():
    cat = motor.leer_catalogo("Repuestos — Carretillas")
    assert len(cat) >= 12, f"el catálogo de carretillas perdió filas: {len(cat)}"
    igual(motor.buscar_repuesto("Bocina de eje", cat)[1]["precio"], Decimal("18.00"), "bocina de eje")
    igual(motor.buscar_repuesto("rueda de carga", cat)[1]["precio"], Decimal("40.00"), "rueda de carga")
    assert motor.buscar_repuesto("Rueda direccional", cat)[0].startswith("Rueda direccional")
    assert motor.buscar_repuesto("Protector soporte de piston", cat) is not None, "sin tilde"
    igual(motor.buscar_repuesto("Tapa protector de resorte", cat)[1]["precio"], Decimal("29.00"),
          "otro nombre del mismo repuesto")
    assert motor.buscar_repuesto("Cilindro que no existe", cat) is None
    try:
        motor.buscar_repuesto("Rueda", cat)
        raise AssertionError("'Rueda' es ambiguo y no debería elegir uno solo")
    except motor.Problema:
        pass


@prueba()
def todas_las_fichas_son_validas():
    fichas = motor.listar_fichas()
    assert fichas, "la biblioteca está vacía"
    for fid in fichas:
        igual(motor.validar_ficha(cargar(f"trabajos/{fid}.json"), fid), [], fid)


@prueba()
def ficha_con_errores_se_detecta():
    ficha = cargar("trabajos/garruchas-coche-panaderia.json")
    ficha["redaccion"] += " {inventado}"
    ficha["condiciones"] = ["Precio de catálogo"]
    ficha["tipo"] = "generales"
    ficha.pop("subtipo")
    errores = motor.validar_ficha(ficha, "garruchas-coche-panaderia")
    contiene(errores, "{inventado}", "marcador sin valor")
    contiene(errores, "términos internos", "texto interno en condiciones")
    contiene(errores, "subtipo", "generales sin subtipo")


# ---------------------------------------------------------------- motor

GARRUCHAS = {"trabajo": "garruchas-coche-panaderia", "cliente": "X", "local": "Y", "unidades": 3,
             "respuestas": {"horno": "si", "montaje": "soldada", "horario": "fuera"}}


@prueba()
def motor_reproduce_el_costeo_real_de_garruchas():
    cot, info = motor.armar(dict(GARRUCHAS), hoy=DIA)
    igual(info["costo"], Decimal("1200"), "costo con SCTR (la hoja original daba 1,150 sin él)")
    igual([p for _, p in info["propuesta"].values()], [Decimal(160), Decimal(170), Decimal(180)],
          "precio por rueda a 60/70/80% (12 ruedas)")
    igual(cot["items"][0]["cant"], 12, "se cotizan 12 ruedas, no 3 juegos")
    igual(cot["subtotal_calculado"], 2040, "subtotal calculado")
    igual(cot["costo_estimado"], 1200, "costo estimado al registro")
    assert cot["items"][0].get("pendiente"), "sin precio_unidad el precio NO se pone solo"
    assert "12 ruedas fenólicas" in cot["items"][0]["descripcion"] and "3 coches" in \
        cot["items"][0]["descripcion"], cot["items"][0]["descripcion"]
    igual(cot["fecha"], "22 de septiembre de 2026", "fecha larga")
    contiene(cot["condiciones"], "fuera del horario", "condición de la respuesta")
    contiene(cot["condiciones"], "se evacúan", "condición de la ficha")

    cot, info = motor.armar(dict(GARRUCHAS, costos={"Juego de 4 ruedas fenólicas con garrucha de acero": 90}), hoy=DIA)
    igual(info["costo"], Decimal("720"), "segundo escenario")
    igual([p for _, p in info["propuesta"].values()], [Decimal(100), Decimal(105), Decimal(110)], "a S/90")
    assert "Costo cambiado para este pedido" in cot["nota_interna"]

    cot, _ = motor.armar(dict(GARRUCHAS, precio_unidad=180, numero="9300-2026"), hoy=DIA)
    igual(generar.calcular(cot)[0], Decimal("2160"), "12 ruedas a S/180")
    assert "+5.9% sobre el calculado" in cot["nota_interna"], cot["nota_interna"]
    igual(validar(cot, final=True, archivar_pedido=True)[0], [], "lista para --final")


@prueba()
def motor_no_calcula_con_costos_que_no_se_saben():
    cot, info = motor.armar(dict(GARRUCHAS, respuestas={"horno": "si", "montaje": "empernada",
                                                        "horario": "dentro"}), hoy=DIA)
    assert info["costo"] is None and info["propuesta"] is None, "la pernería no tiene costo"
    contiene(info["incompleto"], "Pernería", "qué falta")
    assert "subtotal_calculado" not in cot
    assert "pernería nueva" in cot["items"][0]["descripcion"], "la redacción sigue la respuesta"


@prueba()
def motor_exige_el_checklist():
    try:
        motor.armar(dict(GARRUCHAS, respuestas={}), hoy=DIA)
        raise AssertionError("sin respuestas no debería armar")
    except motor.FaltanRespuestas as e:
        igual(len(e.preguntas), 3, "preguntas faltantes")
    cot, info = motor.armar(dict(GARRUCHAS, respuestas={}), borrador=True, hoy=DIA)
    igual(cot["condiciones"][0], "Alcance sujeto a confirmación de: temperatura de trabajo del "
                                 "coche y tipo de montaje de la garrucha.", "borrador")
    assert info["costo"] is None, "el montaje cambia el costo: sin respuesta no hay costo"
    for malo in ({"horno": "tal vez"}, {"color": "rojo"}):
        try:
            motor.armar(dict(GARRUCHAS, respuestas=malo), borrador=True, hoy=DIA)
            raise AssertionError(f"debería rechazar {malo}")
        except motor.Problema:
            pass


@prueba()
def ejemplo_real_de_garruchas_en_borrador():
    cot, info = motor.armar(cargar("trabajos/ejemplos/garruchas-sb-mariategui.json"), borrador=True, hoy=DIA)
    igual(info["costo"], Decimal("1200"), "el horno no cambia el costo registrado")
    igual(cot["condiciones"][0], "Alcance sujeto a confirmación de: temperatura de trabajo del coche.")
    errores, alertas, _ = validar(cot, sin_precios=True)
    igual(errores, [], "sale como propuesta técnica")


@prueba()
def motor_reproduce_las_cotizaciones_emitidas():
    for ejemplo, emitida in (("trabajos/ejemplos/carretillas-hb-panamericana.json", "9197-2026"),
                             ("trabajos/ejemplos/carretillas-tt-pucallpa.json", "9196-2026")):
        cot, info = motor.armar(cargar(ejemplo), hoy=DIA)
        original = cargar(f"cotizaciones-emitidas/{emitida}.json")
        igual(generar.calcular(cot), generar.calcular(original), f"{emitida}: totales")
        igual([it["precio_unit"] for it in cot["items"]],
              [it["precio_unit"] for it in original["items"]], f"{emitida}: precios por línea")
        igual(validar(cot)[0], [], f"{emitida}: lo que arma el motor lo acepta el generador")
        assert {h["numero"] for h in info["historial"]} >= {"9196-2026", "9197-2026"}, "historial"
    cot, _ = motor.armar(cargar("trabajos/ejemplos/carretillas-tt-pucallpa.json"), hoy=DIA)
    assert "Rueda de carga: S/60.00 dado directamente (catálogo: S/40.00" in cot["nota_interna"]
    cot, _ = motor.armar(cargar("trabajos/ejemplos/carretillas-hb-panamericana.json"), hoy=DIA)
    assert "Eje de brazo: S/55.00 de catálogo, POR VALIDAR" in cot["nota_interna"]
    assert "Bajo el piso de S/250.00" in cot["nota_interna"]


@prueba()
def techo_de_precio_con_resultados_reales():
    ficha = cargar("trabajos/mantenimiento-balanza-plataforma.json")
    igual(motor.techo(ficha["referencias"]), (Decimal(5785), Decimal(7690)), "ganada 3032, perdida 8202")
    igual(motor.techo(cargar("trabajos/garruchas-coche-panaderia.json")["referencias"]), None,
          "sin resultados no hay techo")
    igual(len(motor.leer_catalogo("Repuestos — Balanzas")), 4,
          "solo las filas con precio de venta; las de costo de compra se descartan")
    pedido = {"trabajo": "mantenimiento-balanza-plataforma", "cliente": "X", "local": "Y",
              "unidades": 1, "precio_unidad": 5785,
              "respuestas": {"localidad": "lima", "componentes": "marvin", "indicador": "si",
                             "pintura": "no"},
              "lineas": [{"repuesto": "Indicador de peso T-Scale BW", "cant": 1}]}
    cot, _ = motor.armar(pedido, hoy=DIA)
    assert cot["items"][0]["descripcion"].startswith("Mantenimiento de 1 balanza de plataforma:"), \
        "{cantidad} en singular"
    igual(cot["items"][1]["precio_unit"], 1350, "indicador leído del catálogo")
    igual(validar(cot)[0], [], "el generador lo acepta")


@prueba()
def repuesto_fuera_del_catalogo_no_se_inventa():
    base = cargar("trabajos/ejemplos/carretillas-hb-panamericana.json")
    try:
        motor.armar(dict(base, lineas=[{"repuesto": "Cilindro nuevo", "cant": 1}]), hoy=DIA)
        raise AssertionError("un repuesto sin precio no debería armarse")
    except motor.Problema as e:
        assert "no se inventa" in str(e)
    cot, _ = motor.armar(dict(base, lineas=[{"repuesto": "Cilindro nuevo", "cant": 1, "precio": 80}]), hoy=DIA)
    assert "no está en el catálogo" in cot["nota_interna"]


@prueba()
def asistente_crea_una_ficha_que_el_motor_usa():
    import builtins
    import io
    import contextlib
    import nuevo_trabajo
    respuestas = iter([
        "Cambio de burletes de exhibidora", "", "", "", "BURLETES", "exhibidora", "", "", "1",
        "Mano de obra = 150", "Transporte = ?", "",      # costos fijos: uno todavía sin monto
        "Burlete = 80", "",                              # costos por unidad
        "", "",                                          # markup 60-80 por defecto
        "", "¿Es burlete magnético?", "", "",            # comunes + una pregunta propia
        "Cambio de burletes magnéticos en {cantidad} de exhibición.",
        "Los burletes retirados se evacúan del local.", "",
    ])
    tmp = tempfile.mkdtemp()
    trabajos_real, input_real = motor.TRABAJOS, builtins.input
    try:
        motor.TRABAJOS = tmp
        builtins.input = lambda *_: next(respuestas)
        with contextlib.redirect_stdout(io.StringIO()):
            nuevo_trabajo.main()
        ficha = json.load(open(os.path.join(tmp, "cambio-de-burletes-de-exhibidora.json"), encoding="utf-8"))
        igual(ficha["costos_fijos"], {"Mano de obra": 150.0, "Transporte": None}, "el ? queda sin monto")
        igual([p["id"] for p in ficha["preguntas"]], ["localidad", "horario", "es-burlete-magnetico"],
              "preguntas comunes + la propia")
        igual(motor.validar_ficha(ficha, ficha["id"]), [], "la ficha que escribe es válida")
        cot, info = motor.armar({"trabajo": ficha["id"], "cliente": "X", "local": "Y", "unidades": 1,
                                 "respuestas": {"localidad": "lima", "horario": "fuera",
                                                "es-burlete-magnetico": "si"}}, hoy=DIA)
        assert info["propuesta"] is None, "con el transporte sin monto no se propone precio"
        igual(cot["items"][0]["descripcion"],
              "Cambio de burletes magnéticos en 1 exhibidora de exhibición.", "singular")
    finally:
        motor.TRABAJOS, builtins.input = trabajos_real, input_real
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- registro

@prueba()
def laminado_reproduce_cot_2026_001():
    import datetime, re
    import cotizar_laminado as cl
    real = cargar("cotizaciones-emitidas/COT-2026-001.json")
    piezas, esperados = [], []
    for it in real["items"]:
        if it.get("separador"):
            continue
        zona, tipo, a, b = re.search(r"— (.+?): (\w+?)s? de ([\d.]+)m x ([\d.]+)m", it["descripcion"]).groups()
        piezas.append({"zona": zona, "tipo": tipo, "ancho": float(a), "alto": float(b), "cant": it["cant"]})
        esperados.append(it["precio_unit"])
    cot, resumen = cl.armar({"cliente": "X", "local": "Y", "piezas": piezas, "precio_total": 4980},
                            datetime.date(2026, 10, 12))
    items = [i for i in cot["items"] if not i.get("separador")]
    igual([i["precio_unit"] for i in items], esperados, "los 19 precios unitarios de COT-2026-001")
    igual(sum(i["cant"] * i["precio_unit"] for i in items), 4980, "subtotal exacto")
    assert sum(1 for i in cot["items"] if i.get("separador")) == 1, "un separador antes del rollo especial"
    assert any("1.82 m → 10 m" in r for r in resumen), resumen


@prueba()
def laminado_sin_precio_y_piezas_imposibles():
    import datetime
    import cotizar_laminado as cl
    base = {"cliente": "X", "local": "Y"}
    cot, _ = cl.armar(dict(base, piezas=[{"zona": "A", "tipo": "ventana", "ancho": 1, "alto": 1, "cant": 2}]),
                      datetime.date(2026, 10, 12))
    assert all(i.get("pendiente") for i in cot["items"]), "sin precio total: propuesta sin precios"
    for piezas, texto in (([{"zona": "A", "tipo": "espejo", "ancho": 2.0, "alto": 1.9, "cant": 1}], "empalme"),
                          ([{"zona": "A", "tipo": "ventana", "ancho": 97, "alto": 96, "cant": 1}], "metros"),
                          ([], "al menos una pieza")):
        try:
            cl.armar(dict(base, piezas=piezas), datetime.date(2026, 10, 12))
            raise AssertionError(f"debió rechazar: {piezas}")
        except motor.Problema as e:
            assert texto in str(e), e


@prueba()
def balanza_desde_la_app_reproduce_8211():
    import datetime
    pedido = {"cliente": "TOTTUS", "local": "SANTA ANITA", "trabajo": "mantenimiento-balanza-plataforma",
              "unidades": 1, "precio_unidad": 1980,
              "respuestas": {"localidad": "lima", "componentes": "marvin", "indicador": "si", "pintura": "no"},
              "lineas": [{"descripcion": "Cambio de plataforma metálica con plancha estriada", "cant": 1, "precio": 3800},
                         {"repuesto": "Indicador de peso T-Scale BW", "cant": 1, "precio": 1380}]}
    cot, _ = motor.armar(pedido, borrador=True, hoy=datetime.date(2026, 10, 12))
    igual(sum(i["cant"] * i["precio_unit"] for i in cot["items"]), 7160, "mismo subtotal que 8211-2026")
    contiene([cot["items"][0]["descripcion"]], "MAVIN NB2", "la marca elegida va en el texto")
    assert "(" not in cot["items"][2]["descripcion"], "sin notas del catálogo en el texto del cliente"


@prueba()
def tablero_formulas_y_graficos_en_su_lugar():
    import re
    import armar_tablero as tab
    filas, _, g = tab.armar()
    assert all(len(f) == tab.ANCHO for f in filas), "todas las filas con el mismo ancho"
    formulas = [c for f in filas for c in f if isinstance(c, str) and c.startswith("=")]
    assert formulas, "el tablero tiene fórmulas"
    for f in formulas:
        assert f.count("(") == f.count(")"), f"paréntesis sin cerrar: {f[:60]}"
        # SUMIFS no se expande dentro de ARRAYFORMULA: daría el mismo número en toda la columna.
        assert not ("ARRAYFORMULA" in f and "SUMIFS" in f), f"SUMIFS dentro de ARRAYFORMULA: {f[:60]}"
        for hoja, col in re.findall(r"(Cotizaciones|Gastos)!\$([A-Z]+)\$", f):
            limite = "AE" if hoja == "Cotizaciones" else "L"
            assert len(col) < len(limite) or col <= limite, f"columna {hoja}!{col} fuera de la hoja"
    tipos = tab.tipos_de_servicio()
    assert set(tipos) == {"Carretillas", "Balanzas", "Servicios Generales"}, tipos
    # Los gráficos (unas 16 filas cada uno) no se tapan entre sí.
    anclas = [r["addChart"]["chart"]["position"]["overlayPosition"]["anchorCell"]["rowIndex"]
              for r in tab.formato(0) if "addChart" in r]
    assert all(b - a >= 16 for a, b in zip(anclas, anclas[1:])), anclas
    assert g["seguimiento"][1] < g["meses"][0], "el seguimiento no invade la sección de meses"


@prueba()
def registro_agrega_columnas_al_final():
    try:
        import armar_registro as reg
    except ImportError as e:
        raise Omitida(f"falta openpyxl ({e})")
    nombres = [c[0] for c in reg.COLUMNAS]
    igual(nombres[:7], ["N° Cotización", "Fecha cotización", "Cliente", "Local", "Tipo de servicio",
                        "Subtipo", "Descripción breve"], "las primeras columnas no se mueven")
    igual(reg.letra("Precio por unidad"), "Y", "la última columna vieja sigue en Y")
    igual(reg.letra("Costo estimado"), "Z", "las nuevas empiezan en Z")
    igual(reg.letra("Margen % (s/venta)"), "AE", "y terminan en AE")
    assert "otros" in reg.SUBTIPOS.split(","), "subtipos leídos de servicios.json"
    from openpyxl import Workbook
    wb = Workbook()
    ws = reg.hoja_cotizaciones(wb, ["X"])
    assert "AA3" in ws["AC3"].value and "Z3" in ws["AD3"].value, "fórmulas de ajuste y markup"
    reg.hoja_resumen(wb)
    reg.hoja_instrucciones(wb)


@prueba()
def numero_del_dia_mes_dia_y_cuantas_van():
    import datetime
    import procesar_solicitud as ps
    tmp = tempfile.mkdtemp()
    archivo_real = generar.ARCHIVO
    try:
        generar.ARCHIVO = tmp
        dia = datetime.date(2026, 10, 12)
        igual(ps.siguiente_numero(dia), "10121-2026", "la primera del día")
        for n in ("10122-2026", "10124-2026", "10113-2026"):
            open(os.path.join(tmp, n + ".json"), "w").close()
        igual(ps.siguiente_numero(dia), "10125-2026", "sigue a la mayor del día, sin reusar huecos")
        igual(ps.siguiente_numero(datetime.date(2026, 9, 5)), "9051-2026", "mes de un dígito")
        open(os.path.join(tmp, "1011-2026.json"), "w").close()   # 1/01 n° 1: no es del 12/10
        igual(ps.siguiente_numero(datetime.date(2026, 1, 1)), "1012-2026", "enero no choca con octubre")
        for n in range(5, 10):
            open(os.path.join(tmp, f"1012{n}-2026.json"), "w").close()
        try:
            ps.siguiente_numero(dia)
            raise AssertionError("con 9 en el día debe pedir el N° a mano")
        except ps.motor.Problema:
            pass
        assert ps.hoy_lima() == datetime.datetime.now(ps.LIMA).date()
    finally:
        generar.ARCHIVO = archivo_real
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- render (lento)

@prueba(lenta=True)
def render_final_real_sin_tocar_el_archivo():
    if not generar.chromium():
        raise Omitida("no hay Chromium")
    tmp = tempfile.mkdtemp()
    archivo_real = generar.ARCHIVO
    try:
        generar.ARCHIVO = os.path.join(tmp, "archivo")   # nunca pisar el archivo real
        ruta = os.path.join(tmp, "9196-2026.json")
        shutil.copyfile(os.path.join(RAIZ, "cotizaciones-emitidas/9196-2026.json"), ruta)
        argv = sys.argv
        sys.argv = ["generar.py", ruta, tmp, "--final"]
        try:
            generar.main()
        finally:
            sys.argv = argv
        pdf = os.path.join(tmp, "N° 9196-2026 CARRETILLAS TT PUCALLPA MUNAY.pdf")
        subprocess.run(["pdftoppm", "-r", "100", "-png", pdf, os.path.join(tmp, "pag")], check=True)
        from PIL import Image
        im = Image.open(os.path.join(tmp, "pag-1.png")).convert("L")
        w, h = im.size
        px = im.load()
        xs = [x for x in range(w) if any(px[x, y] < 200 for y in range(0, h, 2))]
        ys = [y for y in range(h) if any(px[x, y] < 200 for x in range(0, w, 2))]
        m = [xs[0], w - 1 - xs[-1], ys[0], h - 1 - ys[-1]]
        assert max(m) - min(m) <= 3, f"cotización descentrada en la página, márgenes izq/der/arr/abj: {m}"
        escalones = {sum(1 for y in range(a, a + 40) if px[xs[-1], y] < 100) for a in range(ys[0] + 5, ys[-1] - 40, 80)}
        igual(escalones, {40}, "el borde derecho del marco es continuo")
        for nombre in ("9196-2026.png", "N° 9196-2026 CARRETILLAS TT PUCALLPA MUNAY.pdf",
                       "archivo/9196-2026.jpg", "archivo/9196-2026.json"):
            assert os.path.getsize(os.path.join(tmp, nombre)) > 0, f"falta {nombre}"
    finally:
        generar.ARCHIVO = archivo_real
        shutil.rmtree(tmp, ignore_errors=True)


def main():
    render = "--render" in sys.argv
    fallas = omitidas = corridas = 0
    for f in PRUEBAS:
        if f.lenta and not render:
            continue
        corridas += 1
        try:
            f()
            print(f"  ✓ {f.__name__}")
        except Omitida as e:
            omitidas += 1
            print(f"  - {f.__name__}: omitida ({e})")
        except (Exception, SystemExit):
            fallas += 1
            print(f"  ✗ {f.__name__}")
            traceback.print_exc()
    print(f"\n{corridas - fallas - omitidas} bien, {fallas} con falla, {omitidas} omitidas"
          + ("" if render else "  (con --render se prueba además el PDF real)"))
    sys.exit(1 if fallas else 0)


if __name__ == "__main__":
    main()
